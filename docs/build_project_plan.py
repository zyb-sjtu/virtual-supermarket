"""Generate the 12-week project plan Word document."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("虚拟超市导购智能体十二周项目计划书.docx")
INK = RGBColor(0, 0, 0)
MUTED = RGBColor(83, 91, 98)
NAVY = "24364A"
PALE = "F2F5F8"
BORDER = "D9D9D9"


def set_font(run, size=None, bold=None, color=INK):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_cell(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), BORDER)


def cell_margins(cell, top=90, left=110, bottom=90, right=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Cm(width)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        shade(cell, NAVY)
        border_cell(cell)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, 8.5, True, RGBColor(255, 255, 255))
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = Cm(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border_cell(cells[i])
            cell_margins(cells[i])
            if ri % 2:
                shade(cells[i], PALE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            if i == 0 and len(text) < 10:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_font(r, 8.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def para(doc, text="", style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), 9.5, True)
        set_font(p.add_run(text[len(bold_lead):]), 9.5)
    else:
        set_font(p.add_run(text), 9.5)
    return p


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def week(doc, number, title, work, output, gate):
    heading(doc, f"第 {number} 周 {title}", 2)
    para(doc, f"工作重点  {work}")
    para(doc, f"交付与验收  {output}。验收时检查：{gate}")


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.9)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)
sec.header_distance = Cm(1)
sec.footer_distance = Cm(1)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(9.5)
normal.font.color.rgb = INK
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.paragraph_format.line_spacing = 1.32
normal.paragraph_format.space_after = Pt(6)

title_style = doc.styles["Title"]
title_style.font.name = "Microsoft YaHei"
title_style.font.size = Pt(21)
title_style.font.bold = True
title_style.font.color.rgb = INK
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.paragraph_format.space_after = Pt(12)
title_style.paragraph_format.keep_with_next = True

for name, size, before, after in (("Heading 1", 14, 15, 6), ("Heading 2", 10.5, 10, 4)):
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = INK
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("虚拟超市导购智能体项目计划  |  "), 8, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.add_paragraph("虚拟超市导购智能体十二周项目计划书", style="Title")
p = doc.add_paragraph()
set_font(p.add_run("版本 1.0    编制日期 2026 年 9 月 29 日    计划周期 连续 12 周"), 9, color=MUTED)

heading(doc, "项目目标")
para(doc, "本计划面向一个先在单店验证、再扩展到多店的虚拟超市导购系统。商户可以管理商品和库存，逐步使用图片识别、价格建议和语音修正完成上架；顾客可以通过对话或图片找到商品、获得价格范围内的推荐，并沿店内分区路线前往目标货架。项目以可解释的建议和可追溯的操作记录为原则，涉及商品发布、改价和购买的关键动作均保留人工确认。")
para(doc, "十二周的可交付结果是可供三家模拟或试点超市使用的多租户试点版，以及明确的扩容方案。十二周不承诺完成全国规模运营、真实支付结算或无定位设备的实时室内导航。")

heading(doc, "现有基础与计划边界")
para(doc, "当前仓库已有本机原型：苹果、香蕉两种固定商品，网页购买和补货，SQLite 事务与请求幂等，6 项自动化测试。它验证了最小库存闭环，但尚无自定义上架、角色体系、商品图片处理、价格预测、语音、导购或多店隔离。第一阶段需要把原型改造成可扩展业务系统，而不是把现有功能视为已经完成阶段一。")

heading(doc, "四阶段里程碑")
add_table(doc,
    ["阶段", "周次", "主要成果", "阶段验收"],
    [
        ["基础交易", "1 至 3", "自定义商品上架、分类、购买、库存台账、记录维护和并发保护", "单店完整业务闭环与自动化回归通过"],
        ["价格与双端", "4 至 6", "价格数据管道、建议定价、商品推荐、商户与顾客界面、数据库部署", "价格建议可追溯，角色隔离和部署演练通过"],
        ["双端智能体", "7 至 9", "图片识别上架、语音修正、市场报告、对话导购、图片找货和路线引导", "核心场景评测达到试点门槛，关键动作需确认"],
        ["多店试点", "10 至 12", "租户隔离、多店配置、性能与安全测试、试点验收", "至少三家模拟或试点门店隔离运行"],
    ], [2.4, 2.2, 7.0, 5.2])

heading(doc, "总体技术方案")
para(doc, "建议从现有标准库原型迁移到 Python FastAPI 后端、React 与 TypeScript 双端页面。第一阶段保留本机 SQLite 便于开发和测试，同时用数据访问层及迁移脚本隔离具体数据库；第二阶段部署 PostgreSQL 作为试点环境数据库。图片存入对象存储，数据库只保存媒体地址、标签及审核结果。第三阶段通过模型适配层接入视觉、语音和语言模型，避免业务流程绑定单一供应商。")
add_table(doc,
    ["模块", "职责与关键数据"],
    [
        ["业务服务", "商品 SKU、分类、价格、订单、库存台账、退单、操作审计；库存变动使用事务和幂等键"],
        ["数据与存储", "PostgreSQL 保存结构化数据；对象存储保存图片和语音；价格历史记录来源、时间、适用范围"],
        ["商户端", "上架和改价、库存处理、图片与语音输入、建议审核、市场报告"],
        ["顾客端", "商品检索、范围筛选、对话与图片找货、店内分区路线、明确确认后的购买"],
        ["智能体服务", "只通过受控工具读取商品和地图、生成建议；写操作经规则校验及人工确认后调用业务 API"],
    ], [3.0, 13.8])
para(doc, "核心数据实体包括门店、租户、用户角色、SKU、分类、货架分区、价格历史、库存台账、订单、图片媒体、建议结果和审计事件。第一周就预留 tenant_id 与 store_id 字段及查询边界；第四阶段再启用完整租户隔离策略，减少后期迁移成本。")

heading(doc, "第一阶段 基础交易系统")
week(doc, 1, "需求定稿与数据模型", "梳理商品上架、购买、补货、退单与记录维护流程，定义 SKU、分类、分区和库存台账；确定 API 契约、错误码、角色扩展字段与迁移策略，建立开发环境和持续集成。", "需求清单、数据字典、接口草案、数据库迁移与可运行服务骨架", "现有苹果香蕉数据可无损迁移；库存与订单状态转换有明确规则")
week(doc, 2, "商品自定义上架", "开发商户可用的商品新增、编辑、上下架、图片上传与分类分区维护；加入表单校验、图片大小和格式限制、唯一 SKU 校验及审计日志。", "商品管理页面、商品与分类 API、批量导入样例", "可导入至少 100 个测试 SKU；无效商品不会进入在售列表")
week(doc, 3, "交易库存与记录闭环", "实现顾客购买、取消或退单、补货、库存调整原因记录；订单与库存变化同事务提交，保留请求幂等键并增加并发回归与备份演练。", "单店业务闭环、测试报告、第一阶段演示", "并发购买不超卖；重复提交只记一次；异常回滚后库存和台账一致")

heading(doc, "第二阶段 价格建议与双端分离")
week(doc, 4, "价格数据与角色边界", "建立商户和顾客身份、路由及权限；收集合法可用的历史成交价、手工报价和获授权的市场行情，记录采集时间、地区、规格与来源，并清洗异常值。", "双端页面骨架、角色权限矩阵、价格数据管道", "商户写接口对顾客不可用；每条外部价格均可追溯来源和时间")
week(doc, 5, "建议定价与商品推荐", "先用同品规格历史中位数和近期行情构建可解释基线，再评估时间序列或回归模型；输出价格区间、依据和置信度，支持商户手动修改。顾客端提供分类、价格范围、库存和偏好过滤推荐。", "价格建议服务、手动改价流程、规则加排序的推荐接口", "建议不自动发布；推荐结果满足库存和预算硬约束；给出离线评估与基线对比")
week(doc, 6, "数据库部署与集成验收", "把试点环境迁移到 PostgreSQL，完成连接配置、备份恢复、监控和数据迁移；联调商户端上架改价与顾客端查询购买。", "可重复部署脚本、恢复演练记录、第二阶段验收报告", "新旧数据量和金额核对一致；角色隔离测试通过；可在空环境按文档重建服务")

heading(doc, "第三阶段 双端智能体")
week(doc, 7, "商户图片识别上架", "建立商品图片样本和分区映射，视觉模型提取候选名称、类别、规格与分区；低置信度结果进入人工审核，不直接写入正式商品库。", "图片识别服务、审核界面和标注评测集", "留出测试集分类准确率达到试点目标；所有低置信度样本可人工纠正")
week(doc, 8, "商户定价语音与报告", "将价格建议、库存和市场行情封装为受控工具；商户不同意建议时可手动或语音输入修正，系统展示识别文字及金额供确认。生成注明数据来源、时间和缺失情况的市场报告。", "商户智能体、语音改价流程、市场报告模板", "未经商户确认不发布价格；语音金额单位和小数被回显确认；报告可追溯引用数据")
week(doc, 9, "顾客对话图片与路线", "支持如“推荐 30 元以内的绿叶蔬菜或鱼类”的对话查询、图片找货及多条件推荐；结合门店分区地图生成从选定入口或扫描点到目标货架的路线。建立用户自愿开启的偏好画像与需求预测，提供跨商品比价。", "顾客智能体、图片检索、分区路线与场景评测", "推荐满足明确预算和品类条件；未知商品明确提示；路线仅按已有地图导航；画像可关闭和删除")

heading(doc, "第四阶段 多超市试点")
week(doc, 10, "门店租户隔离", "引入租户管理员、门店成员和顾客权限，所有业务查询强制限定 tenant_id 和 store_id；为 PostgreSQL 配置行级安全策略或等效防线，隔离图片路径、审计记录与缓存键。", "多租户数据迁移、权限配置与隔离测试", "跨店越权查询、更新和媒体访问测试均失败；单店功能回归保持通过")
week(doc, 11, "多店运营与性能", "让每家门店独立配置分类、价格、库存和地图；跨店比价只使用允许公开的商品数据。开展并发、故障恢复、监控告警和提示注入测试。", "多店运营台、压测报告、告警与回滚预案", "预设负载下关键 API 达到响应目标；故障恢复演练可按步骤复现")
week(doc, 12, "试点验收与交付", "准备至少三家模拟或试点门店的数据，组织商户与顾客完整流程验收；修复阻塞问题，冻结版本，交付部署、备份、权限、模型评测和运营手册，并列出下一轮扩容工作。", "多店试点版、验收记录、运维手册和扩容路线图", "三店间数据隔离；核心场景通过；关键缺陷关闭或有明确缓解方案")

heading(doc, "量化验收方法")
para(doc, "下列阈值是试点目标，不是当前已实现性能。第一周固定测试环境、数据量和测量方式；数据不足时保留基线方案，不以模型输出代替真实评测。")
add_table(doc,
    ["指标", "试点目标", "验证方式"],
    [
        ["库存一致性", "并发购买与重复提交用例中，超卖 0 次，重复记账 0 次", "自动化并发测试与数据库台账核对"],
        ["价格建议", "有充分历史数据的品类，相对历史中位数基线的误差改善目标为 10%", "固定留出集；记录样本量、MAE 和偏差；样本不足只展示区间建议"],
        ["商户图片分类", "留出集的一级类目正确率目标不低于 85%", "每类均衡取样，至少 100 张未参与训练的图片"],
        ["顾客查询", "测试问题中，预算及品类等硬约束满足率 100%", "覆盖缺货、超预算、无匹配和歧义条件的场景集"],
        ["多店隔离", "跨租户未授权读写成功数为 0", "API、数据库策略、对象存储和缓存联合测试"],
        ["接口性能", "试点数据量和 50 并发下，主要查询 p95 目标小于 500 ms", "固定机器配置、数据集和压测脚本；结果随环境记录"],
    ], [3.0, 8.0, 5.8])

heading(doc, "团队与协作")
para(doc, "建议以约 4 名全职当量规划：产品与测试合计 1 人、后端与数据 1 人、前端 1 人、AI 与数据分析 1 人；数据库部署和安全评审在第六、十至十二周安排专项支持。人数不足时优先保证第一、二阶段和多店隔离，缩小语音与画像范围。每周展示可运行增量，阶段末由业务代表依据验收记录决定是否进入下一阶段。")

heading(doc, "前置条件与主要风险")
add_table(doc,
    ["事项", "处理方式"],
    [
        ["市场价格数据授权", "仅使用自有交易数据、手工录入或获授权的数据接口；没有可靠行情时标记来源缺失，不生成确定性市场结论。"],
        ["图片与语音质量", "准备有类别、规格和货架位置标注的样本；识别低置信度回到人工编辑，语音金额必须回显确认。"],
        ["室内导航条件", "门店提供分区图、通道和入口点；无定位硬件时仅给出从选定位置开始的静态路线，不宣称实时定位。"],
        ["隐私与推荐偏差", "画像默认关闭，获同意后仅保留必要偏好；允许查看、删除，推荐展示依据并避免敏感属性推断。"],
        ["多租户越权", "从第一周预留租户字段，第四阶段以应用查询约束和数据库策略双重检查，并纳入回归与渗透测试。"],
        ["十二周范围过大", "以三店试点为交付边界；支付、真实货架定位与高可用集群放入后续规划，避免挤压核心质量。"],
    ], [3.3, 13.5])

heading(doc, "交付清单")
para(doc, "项目结束时交付源代码及部署脚本、数据库迁移与备份恢复方案、商户端和顾客端页面、受控智能体与评测集、API 与数据字典、自动化测试和压测报告、三店试点验收记录及运维手册。价格模型与视觉模型还需附训练或调用配置、版本、适用范围、失败回退路径和数据来源说明。")

doc.core_properties.title = "虚拟超市导购智能体十二周项目计划书"
doc.core_properties.subject = "四阶段实施计划与验收方案"
doc.core_properties.author = "项目组"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
